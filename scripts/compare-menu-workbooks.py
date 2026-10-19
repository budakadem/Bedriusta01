"""Report cell-level menu changes between two Bedri Usta workbooks."""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value.strip() if isinstance(value, str) else value


def rows_by_id(sheet):
    rows = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        raw_id = clean(row[0]) if row else ""
        item_id = str(raw_id).zfill(3) if str(raw_id).isdigit() else ""
        if item_id:
            rows[item_id] = [clean(value) for value in row]
    return rows


def compare(before_path: Path, after_path: Path):
    before = load_workbook(before_path, data_only=True, read_only=True)
    after = load_workbook(after_path, data_only=True, read_only=True)
    report = {"beforeSheets": before.sheetnames, "afterSheets": after.sheetnames, "sheets": {}}

    for sheet_name in after.sheetnames:
        if sheet_name not in before.sheetnames:
            report["sheets"][sheet_name] = {"addedSheet": True}
            continue
        before_sheet = before[sheet_name]
        after_sheet = after[sheet_name]
        before_rows = rows_by_id(before_sheet)
        after_rows = rows_by_id(after_sheet)
        headers = [clean(value) for value in next(after_sheet.iter_rows(min_row=3, max_row=3, values_only=True))]
        changes = []
        for item_id in sorted(before_rows.keys() | after_rows.keys()):
            old = before_rows.get(item_id)
            new = after_rows.get(item_id)
            if old is None:
                changes.append({"id": item_id, "kind": "added", "new": new})
                continue
            if new is None:
                changes.append({"id": item_id, "kind": "removed", "old": old})
                continue
            width = max(len(old), len(new), len(headers))
            fields = []
            for index in range(width):
                old_value = old[index] if index < len(old) else ""
                new_value = new[index] if index < len(new) else ""
                if old_value != new_value:
                    fields.append({
                        "column": headers[index] if index < len(headers) and headers[index] else f"Column {index + 1}",
                        "old": old_value,
                        "new": new_value,
                    })
            if fields:
                changes.append({"id": item_id, "kind": "changed", "fields": fields})
        report["sheets"][sheet_name] = {
            "beforeRows": len(before_rows),
            "afterRows": len(after_rows),
            "changes": changes,
        }
    return report


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser()
    parser.add_argument("before", type=Path)
    parser.add_argument("after", type=Path)
    args = parser.parse_args()
    print(json.dumps(compare(args.before, args.after), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
